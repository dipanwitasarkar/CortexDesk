import logging
from typing import Dict, Any, Optional, List
from pathlib import Path
import json
import os
from datetime import datetime
from app.core.database import get_async_db
from app.models.document import Document, DocumentType
from app.services.llm_service import llm_service
from app.core.qdrant import qdrant_manager
from sqlalchemy import select, delete
import re


class DocumentService:
    """Service for document processing, chunking, and embedding"""
    
    def __init__(self):
        self.logger = logging.getLogger("document_service")
        self.chunk_size = 500  # characters per chunk
        self.chunk_overlap = 50  # characters overlap between chunks
    
    async def create_document(
        self,
        user_id: int,
        title: str,
        file_type: str,
        content: Optional[str] = None,
        file_path: Optional[str] = None,
        metadata: Optional[Dict[str, Any]] = None
    ) -> int:
        """Create a new document entry"""
        async for db in get_async_db():
            try:
                document_type = DocumentType(file_type.lower())
            except ValueError:
                document_type = DocumentType.TXT
            
            document = Document(
                user_id=user_id,
                title=title,
                file_type=document_type,
                content=content,
                file_path=file_path,
                meta_data=json.dumps(metadata) if metadata else None,
                embedding_status="pending"
            )
            db.add(document)
            await db.commit()
            await db.refresh(document)
            
            self.logger.info(f"Created document {document.id}: {title}")
            return document.id
    
    async def process_document(
        self,
        document_id: int,
        content: str
    ) -> Dict[str, Any]:
        """Process document: chunking and embedding"""
        async for db in get_async_db():
            # Get document
            result = await db.execute(
                select(Document).where(Document.id == document_id)
            )
            document = result.scalar_one_or_none()
            
            if not document:
                raise ValueError(f"Document {document_id} not found")
            
            # Update status to processing
            document.embedding_status = "processing"
            await db.commit()
            
            try:
                # Chunk the content
                chunks = self._chunk_text(content)
                chunk_count = len(chunks)
                
                # Try to generate embeddings
                try:
                    embeddings = []
                    for i, chunk in enumerate(chunks):
                        try:
                            embedding = await llm_service.generate_embedding(chunk)
                            embeddings.append({
                                "chunk_id": f"{document_id}_chunk_{i}",
                                "text": chunk,
                                "vector": embedding,
                                "metadata": {
                                    "document_id": document_id,
                                    "chunk_index": i,
                                    "user_id": document.user_id,
                                    "title": document.title
                                }
                            })
                        except Exception as e:
                            self.logger.error(f"Failed to generate embedding for chunk {i}: {e}")
                    
                    self.logger.info(f"Generated {len(embeddings)} embeddings")
                    self.logger.info(f"Sample embedding length: {len(embeddings[0]['vector']) if embeddings else 'none'}")
                    self.logger.info(f"Sample embedding type: {type(embeddings[0]['vector']) if embeddings else 'none'}")
                    self.logger.info(f"Sample embedding first 5 values: {embeddings[0]['vector'][:5] if embeddings else 'none'}")
                    
                    # Store in Qdrant if embeddings were generated
                    if embeddings:
                        await qdrant_manager.create_collection("documents", vector_size=384)
                        points = [
                            {
                                "id": emb["chunk_id"],
                                "vector": emb["vector"],
                                "payload": emb["metadata"]
                            }
                            for emb in embeddings
                        ]
                        self.logger.info(f"Inserting {len(points)} points into Qdrant")
                        self.logger.info(f"Sample point: {points[0]}")
                        await qdrant_manager.insert_points("documents", points)
                except Exception as e:
                    self.logger.error(f"Embedding generation/storage failed: {e}")
                    # Mark as failed
                    embeddings = []
                
                # Update document status
                document.chunk_count = chunk_count
                document.embedding_status = "completed" if embeddings else "failed"
                await db.commit()
                
                self.logger.info(f"Processed document {document_id}: {chunk_count} chunks, {len(embeddings)} embedded")
                
                return {
                    "document_id": document_id,
                    "chunk_count": chunk_count,
                    "embedding_status": "completed" if embeddings else "failed",
                    "chunks_processed": len(embeddings),
                    "note": None if embeddings else "Embedding generation failed"
                }
                
            except Exception as e:
                document.embedding_status = "failed"
                await db.commit()
                self.logger.error(f"Failed to process document {document_id}: {e}")
                raise
    
    def _chunk_text(self, text: str) -> List[str]:
        """Split text into chunks with overlap"""
        if not text:
            return []
        
        chunks = []
        start = 0
        
        while start < len(text):
            end = start + self.chunk_size
            
            # Try to break at word boundary
            if end < len(text):
                # Find last space before end
                last_space = text.rfind(' ', start, end)
                if last_space > start:
                    end = last_space
            
            chunk = text[start:end].strip()
            if chunk:
                chunks.append(chunk)
            
            start = end - self.chunk_overlap if end > self.chunk_overlap else end
        
        return chunks
    
    async def search_documents(
        self,
        user_id: int,
        query: str,
        limit: int = 5
    ) -> List[Dict[str, Any]]:
        """Search documents using semantic search"""
        try:
            # Generate query embedding
            query_embedding = await llm_service.generate_embedding(query)
            
            # Search Qdrant
            results = await qdrant_manager.search(
                "documents",
                query_embedding,
                limit=limit,
                filter_dict={"user_id": user_id}
            )
            
            return results
        except Exception as e:
            self.logger.error(f"Failed to search documents: {e}")
            return []
    
    async def get_document(self, document_id: int) -> Optional[Dict[str, Any]]:
        """Get document by ID"""
        async for db in get_async_db():
            result = await db.execute(
                select(Document).where(Document.id == document_id)
            )
            document = result.scalar_one_or_none()
            
            if not document:
                return None
            
            return {
                "id": document.id,
                "user_id": document.user_id,
                "title": document.title,
                "file_type": document.file_type.value,
                "content": document.content,
                "file_path": document.file_path,
                "chunk_count": document.chunk_count,
                "embedding_status": document.embedding_status,
                "meta_data": json.loads(document.meta_data) if document.meta_data else None,
                "created_at": document.created_at.isoformat(),
                "updated_at": document.updated_at.isoformat()
            }
    
    async def list_documents(
        self,
        user_id: int,
        limit: int = 50,
        offset: int = 0
    ) -> List[Dict[str, Any]]:
        """List documents for a user"""
        async for db in get_async_db():
            result = await db.execute(
                select(Document)
                .where(Document.user_id == user_id)
                .order_by(Document.created_at.desc())
                .limit(limit)
                .offset(offset)
            )
            documents = result.scalars().all()
            
            return [
                {
                    "id": doc.id,
                    "user_id": doc.user_id,
                    "title": doc.title,
                    "file_type": doc.file_type.value,
                    "content": doc.content,
                    "file_path": doc.file_path,
                    "chunk_count": doc.chunk_count,
                    "embedding_status": doc.embedding_status,
                    "meta_data": json.loads(doc.meta_data) if doc.meta_data else None,
                    "created_at": doc.created_at.isoformat(),
                    "updated_at": doc.updated_at.isoformat()
                }
                for doc in documents
            ]
    
    async def delete_document(self, document_id: int) -> bool:
        """Delete a document and its embeddings"""
        async for db in get_async_db():
            # Get document
            result = await db.execute(
                select(Document).where(Document.id == document_id)
            )
            document = result.scalar_one_or_none()
            
            if not document:
                return False
            
            # Delete embeddings from Qdrant
            try:
                chunk_ids = [f"{document_id}_chunk_{i}" for i in range(document.chunk_count)]
                await qdrant_manager.delete_points("documents", chunk_ids)
            except Exception as e:
                self.logger.warning(f"Failed to delete embeddings: {e}")
            
            # Delete document from database
            await db.execute(delete(Document).where(Document.id == document_id))
            await db.commit()
            
            self.logger.info(f"Deleted document {document_id}")
            return True
    
    async def process_text_content(
        self,
        user_id: int,
        title: str,
        text_content: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> int:
        """Create and process a text document in one step"""
        # Create document
        document_id = await self.create_document(
            user_id=user_id,
            title=title,
            file_type="txt",
            content=text_content,
            metadata=metadata
        )
        
        # Process document (chunking + embedding)
        await self.process_document(document_id, text_content)
        
        return document_id
    
    async def extract_text_from_file(self, file_path: str, filename: str) -> str:
        """Extract text from document files (PDF, DOCX, PPTX)"""
        ext = os.path.splitext(filename)[1].lower()
        
        try:
            if ext == '.pdf':
                return await self._extract_from_pdf(file_path)
            elif ext in ['.docx', '.doc']:
                return await self._extract_from_docx(file_path)
            elif ext in ['.pptx', '.ppt']:
                return await self._extract_from_pptx(file_path)
            else:
                # Fallback: try to read as text
                with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                    return f.read()
        except Exception as e:
            self.logger.error(f"Failed to extract text from {filename}: {e}")
            raise ValueError(f"Failed to extract text from file: {str(e)}")
    
    async def _extract_from_pdf(self, file_path: str) -> str:
        """Extract text from PDF file"""
        try:
            import PyPDF2
            text = ""
            with open(file_path, 'rb') as file:
                reader = PyPDF2.PdfReader(file)
                for page in reader.pages:
                    text += page.extract_text() + "\n"
            return text
        except ImportError:
            self.logger.warning("PyPDF2 not installed, trying alternative method")
            # Fallback: try to read as text
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            self.logger.error(f"Failed to extract from PDF: {e}")
            raise
    
    async def _extract_from_docx(self, file_path: str) -> str:
        """Extract text from DOCX file"""
        try:
            from docx import Document
            doc = Document(file_path)
            text = ""
            for paragraph in doc.paragraphs:
                text += paragraph.text + "\n"
            return text
        except ImportError:
            self.logger.warning("python-docx not installed, trying alternative method")
            # Fallback: try to read as text
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            self.logger.error(f"Failed to extract from DOCX: {e}")
            raise
    
    async def _extract_from_pptx(self, file_path: str) -> str:
        """Extract text from PPTX file"""
        try:
            from pptx import Presentation
            prs = Presentation(file_path)
            text = ""
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
            return text
        except ImportError:
            self.logger.warning("python-pptx not installed, trying alternative method")
            # Fallback: try to read as text
            with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
                return f.read()
        except Exception as e:
            self.logger.error(f"Failed to extract from PPTX: {e}")
            raise
    
    async def extract_text_from_image(self, file_path: str) -> str:
        """Extract text from image using OCR"""
        try:
            import pytesseract
            from PIL import Image
            
            # Open image and extract text
            image = Image.open(file_path)
            text = pytesseract.image_to_string(image)
            return text
        except ImportError:
            self.logger.warning("pytesseract or PIL not installed, returning placeholder")
            return "[Image file uploaded - OCR not available. Install pytesseract and pillow to extract text from images.]"
        except Exception as e:
            self.logger.error(f"Failed to extract text from image: {e}")
            return f"[Failed to extract text from image: {str(e)}]"


# Global instance
document_service = DocumentService()